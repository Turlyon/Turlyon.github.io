#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""从 管网运营场景页面清单.xlsx 生成 12 份模块 spec markdown（供子 agent 生成原型使用）"""
import openpyxl
import os
import re

XLSX = "/Users/liuguanang/WorkBuddy/2026-10-12-10-42-05/管网运营场景页面清单.xlsx"
OUT_DIR = "/Users/liuguanang/WorkBuddy/2026-10-12-10-42-05/原型/spec"
os.makedirs(OUT_DIR, exist_ok=True)

wb = openpyxl.load_workbook(XLSX)

# ---- 读模块总览 ----
ws0 = wb["模块总览"]
overview = {}  # name -> dict
order = []
for r in range(3, ws0.max_row + 1):
    name = ws0.cell(row=r, column=2).value
    if not name or name == "合计":
        continue
    overview[name] = {
        "clause": ws0.cell(row=r, column=3).value or "",
        "position": ws0.cell(row=r, column=4).value or "",
        "pages": ws0.cell(row=r, column=5).value or 0,
        "capability": ws0.cell(row=r, column=6).value or "",
        "kpi": ws0.cell(row=r, column=7).value or "",
        "priority": ws0.cell(row=r, column=8).value or "",
    }
    order.append(name)

# ---- 读页面清单 ----
ws = wb["页面清单"]
pages = {name: [] for name in order}
for r in range(2, ws.max_row + 1):
    mod = ws.cell(row=r, column=2).value
    if mod not in pages:
        continue
    pages[mod].append({
        "page": ws.cell(row=r, column=3).value or "",
        "terminal": ws.cell(row=r, column=4).value or "",
        "roles": ws.cell(row=r, column=5).value or "",
        "funcs": ws.cell(row=r, column=6).value or "",
        "fields": ws.cell(row=r, column=7).value or "",
        "inter": ws.cell(row=r, column=8).value or "",
        "clause": ws.cell(row=r, column=9).value or "",
    })

# ---- 读调研补充（Sheet4 Part2 分模块调研发现） ----
ws3 = wb["智能体优先级与调研补充"]
survey = {}  # 智能体名(近似) -> (发现, 待明确)
cur_section = None
for r in range(1, ws3.max_row + 1):
    a = ws3.cell(row=r, column=1).value
    b = ws3.cell(row=r, column=2).value
    c = ws3.cell(row=r, column=3).value
    d = ws3.cell(row=r, column=4).value
    if isinstance(a, str) and a.startswith("二、"):
        cur_section = "survey"
        continue
    if isinstance(a, str) and a.startswith("三、"):
        cur_section = None
    if cur_section == "survey" and isinstance(a, int) and b:
        survey[str(b)] = (c or "", d or "")

SLUGS = {
    "管网生命线一张图": ("m01", "map"),
    "管网健康等级评估助手": ("m02", "health"),
    "管道巡检智能排班助手": ("m03", "schedule"),
    "管道保护AI图像稽核助手": ("m04", "patrol-audit"),
    "管道保护语音校验助手": ("m05", "voice"),
    "应急抢修资源推荐助手": ("m06", "rescue"),
    "关阀方案推荐智能体": ("m07", "valve"),
    "智能导航助手": ("m08", "nav"),
    "关键设备AI图像稽核助手": ("m09", "device-audit"),
    "关键设备管理": ("m10", "device"),
    "员工工作效能统计与提升": ("m11", "efficiency"),
    "气源设计方案比选推荐助手": ("m12", "compare"),
    "燃气管网工程天然气气源设计方案比选推荐助手": ("m12", "compare"),
}

# 调研表智能体名 -> 模块名 的映射（模糊匹配）
def find_survey(mod_name):
    for key, val in survey.items():
        k = key.replace(" ", "")
        m = mod_name.replace(" ", "")
        if k in m or m in k:
            return val
    return None

def view_id(page_name):
    s = re.sub(r"[·•／/（）()\s]+", "-", page_name)
    s = re.sub(r"[^\w\u4e00-\u9fff-]", "", s)
    return "v-" + s[:40]

for idx, mod in enumerate(order, 1):
    o = overview[mod]
    code, slug = SLUGS[mod]
    lines = []
    lines.append(f"# {mod}（M{idx:02d}）模块规格")
    lines.append("")
    lines.append(f"> 原型文件：`modules/{code}-{slug}.html` ｜ 对应页面清单「{mod}」")
    lines.append("")
    lines.append("## 模块信息")
    lines.append("")
    lines.append(f"- **对应条款**：{o['clause']}")
    lines.append(f"- **页面数**：{o['pages']}")
    lines.append(f"- **优先级/节奏**：{o['priority']}")
    lines.append(f"- **考核指标**：{o['kpi']}")
    lines.append(f"- **建设定位**：{o['position']}")
    lines.append("")
    sv = find_survey(mod)
    if sv:
        lines.append("## 调研关键发现与待明确")
        lines.append("")
        lines.append(f"- **调研关键发现**：{sv[0]}")
        lines.append(f"- **待明确事项**：{sv[1]}")
        lines.append("")
    lines.append("## 页面清单（每页 = 原型内一个 view，view id 已定义）")
    lines.append("")
    for j, p in enumerate(pages[mod], 1):
        vid = view_id(p["page"])
        lines.append(f"### 页面{j}：{p['page']}")
        lines.append("")
        lines.append(f"- **view id**：`{vid}`")
        lines.append(f"- **终端**：{p['terminal']}")
        lines.append(f"- **使用角色**：{p['roles']}")
        lines.append(f"- **条款**：{p['clause']}")
        lines.append("")
        lines.append(f"**功能模块（页面承载功能）**：{p['funcs']}")
        lines.append("")
        lines.append(f"**关键字段/数据项**：{p['fields']}")
        lines.append("")
        lines.append(f"**交互逻辑**：{p['inter']}")
        lines.append("")
        lines.append("---")
        lines.append("")
    # 原型实现建议
    has_app = any(("APP" in p["terminal"] or "移动" in p["terminal"] or "手机" in p["terminal"]) for p in pages[mod])
    tips = ["顶部 module-top：模块名 + 优先级徽标 + segmented control 切换各页面 view",
            "地图/GIS 一律用 .map-canvas CSS 示意（网格底+彩色管线+点位），禁外网地图"]
    if has_app:
        tips.append("含外勤APP/移动端的页面必须用 .phone-frame 手机框呈现")
    tips.append("演示数据用真实感中文数据（管段GD-XXXX、工单号、脱敏姓名），基数与调研一致")
    lines.append("## 原型实现要求")
    lines.append("")
    for t in tips:
        lines.append(f"- {t}")
    lines.append("")
    fname = f"{code}-{slug}.md"
    with open(os.path.join(OUT_DIR, fname), "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print(f"✓ {fname} ({o['pages']}页)")

print("\n全部完成")
